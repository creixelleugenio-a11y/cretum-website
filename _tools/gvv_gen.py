import json, sys, os
from openpyxl import load_workbook
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__))); from gvv_agg import aggregate
SECTOR_FULL=json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)),'sector_full.json')))
REGION={'US':'Norteamérica','MM':'América Latina','LN':'Europa','CN':'Norteamérica','SS':'Europa','GY':'Europa','MXN':'América Latina','NA':'Europa'}
def strat_eq(e):
    e=str(e or ''); return e[7:] if e.startswith('CRETUM ') else e
def num(v): 
    try: return float(v)
    except: return 0
def excel_holdings(eq):
    H=[]
    for r in range(3,eq.max_row+1):
        t=eq.cell(r,1).value
        if t is None: continue
        titulos=num(eq.cell(r,5).value); value=num(eq.cell(r,6).value); pl=num(eq.cell(r,8).value)
        sec=eq.cell(r,4).value; sec='nan' if sec is None else SECTOR_FULL.get(str(sec),str(sec))
        mkt=str(eq.cell(r,3).value)
        H.append({'ticker':str(t),'company':str(eq.cell(r,2).value),'mkt':mkt,'sector':sec,
            'region':REGION.get(mkt,'Europa'),'ac':str(eq.cell(r,11).value),'strat':strat_eq(eq.cell(r,12).value),
            'titulos':titulos,'spot':(value/titulos if titulos else value),'value':value,'pl':pl,
            'plpct':(pl/(value-pl)*100 if (value-pl) else 0)})
    return H
def fdate(v):
    if v is None: return ''
    try: return v.strftime('%d/%m/%Y')
    except: return str(v)
def excel_options(op):
    O=[]
    for r in range(2,op.max_row+1):
        t=op.cell(r,1).value
        if t is None: continue
        O.append({'ticker':str(t),'mkt':str(op.cell(r,3).value),'ultimo':num(op.cell(r,4).value),
            'strat':str(op.cell(r,5).value),'estado':str(op.cell(r,6).value),'cv':str(op.cell(r,7).value),
            'pc':str(op.cell(r,8).value),'strike':num(op.cell(r,9).value),'titulos':num(op.cell(r,11).value),
            'prima':num(op.cell(r,12).value),'be':num(op.cell(r,13).value),'primas':num(op.cell(r,14).value),
            'pl':num(op.cell(r,17).value),'rend':num(op.cell(r,18).value)*100,'venc':fdate(op.cell(r,21).value),
            'dias':num(op.cell(r,22).value)})
    return O
def agg_opt(O):
    strat={}
    for o in O:
        d=strat.setdefault(o['strat'],{'n':0,'pl':0,'primas':0}); d['n']+=1; d['pl']+=o['pl']; d['primas']+=o['primas']
    return {'opt_n':len(O),'opt_strat':sorted([{'label':k,'n':d['n'],'pl':d['pl'],'primas':d['primas']} for k,d in strat.items()],key=lambda x:-x['n']),
        'opt_pl':sum(o['pl'] for o in O),'opt_primas':sum(o['primas'] for o in O),
        'opt_itm':sum(1 for o in O if o['estado']=='ITM'),'opt_otm':sum(1 for o in O if o['estado']=='OTM'),
        'opt_venta':sum(1 for o in O if o['cv']=='Venta'),'opt_compra':sum(1 for o in O if o['cv']=='Compra'),
        'opt_und':len(set(o['ticker'] for o in O)),'options':O}
def build(path):
    wb=load_workbook(path,data_only=True)
    H=excel_holdings(wb['Equities']); O=excel_options(wb['Opciones'])
    D=aggregate(H); D['holdings']=H; D.update(agg_opt(O))
    return D
if __name__=='__main__':
    D=build('/Users/air/Downloads/Actualización Tickers GVV análisis Claude (1).xlsx')
    print("total: ${:,.0f}".format(D['total']),"| total_pl: ${:,.0f}".format(D['total_pl']),"| n_lines:",D['n_lines'],"| n_companies:",D['n_companies'])
    print("cash_v: ${:,.0f}".format(D['cash_v']),"| priv_v: ${:,.0f}".format(D['priv_v']),"| pub_v: ${:,.0f}".format(D['pub_v']))
    print("asset_class:",[(a['label'],round(a['pct'],1)) for a in D['asset_class']])
    print("geography:",[(g['label'],round(g['pct'],1)) for g in D['geography']])
    print("strategy:",[(s['label'],s['count']) for s in D['strategy']])
    print("top3:",[(h['company'],round(h['pct'],1)) for h in D['top10'][:3]])
    print("privates:",len(D['privates']),"| priv_total ${:,.0f}".format(D['priv_total']))
    print("opt_n:",D['opt_n'],"| opt_strat:",[(s['label'],s['n']) for s in D['opt_strat']],"| opt_pl ${:,.0f}".format(D['opt_pl']))
    print("\nej holding:",json.dumps(D['holdings'][0],ensure_ascii=False))
    print("ej option:",json.dumps(D['options'][0],ensure_ascii=False))
    json.dump(D,open('/tmp/gvv_newD.json','w'),ensure_ascii=False)
